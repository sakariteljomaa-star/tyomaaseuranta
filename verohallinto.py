"""
Verohallinnon rakentamisilmoitukset — päätoteuttajan tiedonantovelvollisuus.

VAIHE 2: Työntekijäilmoitus (ketkä työskentelivät työmaalla kuukauden aikana).
Kokoaa tuntikirjan (ali_tunnit) + tekijöiden viranomaistiedot kuukausittaiseksi
raportiksi: Excel + CSV (OmaVeroon syötettäväksi tai kirjanpitäjälle).

HUOM Ilmoitin.fi-tiedostomuodosta:
Verohallinnon kiinteämittainen tiedostomuoto (VSRAKYHT/VSTYONTE) muuttuu
versioittain. Tämä moduuli tuottaa LUOTETTAVAN Excel/CSV-listan kaikilla
vaadituilla kentillä. Automaattinen Ilmoitin.fi-tiedosto lisätään kun
voimassa oleva tietuekuvaus on varmistettu — älä lähetä kiinteämittaista
tiedostoa tarkistamatta formaattia.
"""

import io
from datetime import date, timedelta
import pandas as pd

PAIVA_AVAIN = ["ma", "ti", "ke", "to", "pe", "la", "su"]


def _viikon_paivat(vuosi: int, viikko: int):
    ma = date.fromisocalendar(vuosi, viikko, 1)
    return [ma + timedelta(days=i) for i in range(7)]


def _tekijan_paivat_kuukaudessa(rivi: dict, vuosi: int, kk: int) -> list:
    """Palauttaa päivämäärät (date) joina tekijällä oli tunteja kyseisessä kuussa."""
    v = int(rivi.get("vuosi", vuosi))
    vk = int(rivi.get("viikko", 0))
    if not vk:
        return []
    paivat = _viikon_paivat(v, vk)
    tulos = []
    for pk, pvm in zip(PAIVA_AVAIN, paivat):
        tunnit = rivi.get(pk, 0) or 0
        if tunnit and pvm.year == vuosi and pvm.month == kk:
            tulos.append(pvm)
    return tulos


def kokoa_tyontekijailmoitus(ali_rivit: list, tyontekijat: list,
                             projekti_info: dict, vuosi: int, kk: int) -> pd.DataFrame:
    """
    Palauttaa DataFramen työmaalla kuukauden aikana työskennelleistä.
    Yhdistää tuntidatan (läsnäolopäivät) tekijöiden viranomaistietoihin.
    """
    # Hae tekijän viranomaistiedot nimen perusteella
    tiedot = {t["nimi"]: t for t in tyontekijat}

    # Kerää per tekijä läsnäolopäivät kuukaudessa
    per_tekija = {}
    for rivi in ali_rivit:
        paivat = _tekijan_paivat_kuukaudessa(rivi, vuosi, kk)
        if not paivat:
            continue
        nimi = rivi.get("nimi", "")
        per_tekija.setdefault(nimi, []).extend(paivat)

    rivit = []
    for nimi, paivat in per_tekija.items():
        t = tiedot.get(nimi, {})
        alku = min(paivat)
        loppu = max(paivat)
        rivit.append({
            "Nimi":              nimi,
            "Veronumero":        t.get("veronumero", ""),
            "Syntymäaika":       t.get("syntymaaika", ""),
            "Työnantaja":        t.get("yritys", ""),
            "Y-tunnus":          t.get("y_tunnus", ""),
            "Työsuhteen laji":   t.get("tyosuhde", ""),
            "Kansalaisuus":      t.get("kansalaisuus", ""),
            "Lähettävä yritys":  t.get("laheva_yritys", ""),
            "Työmaalla alku":    alku.strftime("%d.%m.%Y"),
            "Työmaalla loppu":   loppu.strftime("%d.%m.%Y"),
            "Työpäiviä":         len(set(paivat)),
            "Osoite":            t.get("osoite", ""),
            "Puhelin":           t.get("puhelin", ""),
            "_puuttuu":          not (t.get("veronumero") and t.get("syntymaaika") and t.get("y_tunnus")),
        })
    df = pd.DataFrame(rivit)
    return df.sort_values("Nimi").reset_index(drop=True) if not df.empty else df


def tyontekijailmoitus_excel(df: pd.DataFrame, projekti_info: dict,
                             vuosi: int, kk: int) -> bytes:
    """Tuottaa tulostettavan/lähetettävän Excel-ilmoituksen."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = f"Tyontekijat {kk}-{vuosi}"
    THIN = Side(style="thin", color="AAAAAA")
    B = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # Otsikko
    ws.merge_cells("A1:M1")
    c = ws.cell(1, 1, "TYÖNTEKIJÄILMOITUS — Verohallinto (rakentamisen tiedonantovelvollisuus)")
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F3864")

    # Työmaan tiedot
    ws.merge_cells("A2:M2")
    tyomaa = (f"Työmaa: {projekti_info.get('nimi','')}  |  "
              f"Osoite: {projekti_info.get('tyomaa_osoite','')}  |  "
              f"Työmaa-avain: {projekti_info.get('tyomaa_avain','')}")
    ws.cell(2, 1, tyomaa).font = Font(size=10)
    ws.merge_cells("A3:M3")
    paatot = (f"Päätoteuttaja: {projekti_info.get('paatoteuttaja','')} "
              f"({projekti_info.get('paatot_ytunnus','')})  |  "
              f"Kausi: {kk:02d}/{vuosi}")
    ws.cell(3, 1, paatot).font = Font(size=10)

    # Sarakeotsikot
    otsikot = ["Nimi","Veronumero","Syntymäaika","Työnantaja","Y-tunnus",
               "Työsuhteen laji","Kansalaisuus","Lähettävä yritys",
               "Työmaalla alku","Työmaalla loppu","Työpäiviä","Osoite","Puhelin"]
    r0 = 5
    for j, ot in enumerate(otsikot, 1):
        cc = ws.cell(r0, j, ot)
        cc.font = Font(bold=True, color="FFFFFF")
        cc.fill = PatternFill("solid", fgColor="2E75B6")
        cc.alignment = Alignment(wrap_text=True, vertical="center")
        cc.border = B
    ws.row_dimensions[r0].height = 30

    # Data
    for i, (_, rivi) in enumerate(df.iterrows(), r0 + 1):
        puuttuu = rivi.get("_puuttuu", False)
        for j, ot in enumerate(otsikot, 1):
            cc = ws.cell(i, j, rivi.get(ot, ""))
            cc.border = B
            cc.font = Font(size=10)
            if puuttuu and ot in ("Veronumero","Syntymäaika","Y-tunnus") and not rivi.get(ot):
                cc.fill = PatternFill("solid", fgColor="FFCDD2")  # punainen = puuttuu

    # Leveydet
    for col, w in zip("ABCDEFGHIJKLM", [22,14,12,24,12,18,12,20,13,13,9,24,14]):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tyontekijailmoitus_csv(df: pd.DataFrame) -> bytes:
    """CSV (puolipiste-eroteltu, UTF-8 BOM) OmaVeroon/kirjanpitäjälle."""
    nayta = df.drop(columns=["_puuttuu"], errors="ignore")
    return nayta.to_csv(index=False, sep=";").encode("utf-8-sig")


# ── VAIHE 3: Urakkailmoitus (maksetut summat aliurakoitsijoille) ───────────────

def poimi_toimittaja(selite: str) -> str:
    """Yrittää poimia toimittajan nimen ostolaskun selitteestä (best effort)."""
    import re
    s = str(selite or "").strip()
    # Katkaise ennen ", lasku" / " – lasku" / " - lasku"
    osat = re.split(r"\s*[,–\-]\s*lasku", s, flags=re.IGNORECASE)
    nimi = osat[0].strip(" –-,")
    # Jos selite ei sisällä nimeä (esim. "Työ (vko 20)", "Aloitus pvm...", "Lasku 77227 ...")
    if re.match(r"^(aloitus|lopetus|työ|vko|viikko|lasku)\b", nimi, re.IGNORECASE):
        return ""
    return nimi


def kokoa_urakkailmoitus(df_ostot, yritysrekisteri: dict,
                         vuosi: int, kk: int) -> pd.DataFrame:
    """
    Kokoaa aliurakoitsijoille kaudella maksetut summat.
    df_ostot: kustannusseurannan ostolasku-DataFrame (sis. kululaji, summa, pvm, selite).
    yritysrekisteri: {toimittaja: {"y_tunnus":..., "yhteyshenkilo":...}}
    """
    if df_ostot is None or df_ostot.empty:
        return pd.DataFrame()

    df = df_ostot[df_ostot["kululaji"] == "Aliurakoitsijat"].copy()
    if df.empty:
        return pd.DataFrame()

    df["pvm"] = pd.to_datetime(df["pvm"], errors="coerce")
    df = df[(df["pvm"].dt.year == vuosi) & (df["pvm"].dt.month == kk)]
    if df.empty:
        return pd.DataFrame()

    df["toimittaja"] = df["selite"].apply(poimi_toimittaja)

    rivit = []
    for toimittaja, osa in df.groupby("toimittaja"):
        reko = yritysrekisteri.get(toimittaja, {})
        rivit.append({
            "Toimittaja":     toimittaja or "(tuntematon — täydennä)",
            "Y-tunnus":       reko.get("y_tunnus", ""),
            "Yhteyshenkilö":  reko.get("yhteyshenkilo", ""),
            "Summa (alv 0)":  round(float(osa["summa"].sum()), 2),
            "Laskuja":        int(osa["lasku"].nunique()) if "lasku" in osa else len(osa),
        })
    df_out = pd.DataFrame(rivit).sort_values("Toimittaja").reset_index(drop=True)
    return df_out


def urakkailmoitus_excel(df: pd.DataFrame, projekti_info: dict,
                         vuosi: int, kk: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = f"Urakat {kk}-{vuosi}"
    THIN = Side(style="thin", color="AAAAAA")
    B = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    ws.merge_cells("A1:E1")
    c = ws.cell(1, 1, "URAKKAILMOITUS — Verohallinto (rakentamisen tiedonantovelvollisuus)")
    c.font = Font(bold=True, size=13, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
    ws.merge_cells("A2:E2")
    ws.cell(2, 1, f"Työmaa: {projekti_info.get('nimi','')} | "
                  f"Tilaaja: {projekti_info.get('paatoteuttaja','')} "
                  f"({projekti_info.get('paatot_ytunnus','')}) | Kausi: {kk:02d}/{vuosi}").font = Font(size=10)

    otsikot = ["Toimittaja","Y-tunnus","Yhteyshenkilö","Summa (alv 0)","Laskuja"]
    r0 = 4
    for j, ot in enumerate(otsikot, 1):
        cc = ws.cell(r0, j, ot); cc.font = Font(bold=True, color="FFFFFF")
        cc.fill = PatternFill("solid", fgColor="2E75B6"); cc.border = B
    for i, (_, rv) in enumerate(df.iterrows(), r0 + 1):
        for j, ot in enumerate(otsikot, 1):
            cc = ws.cell(i, j, rv.get(ot, "")); cc.border = B; cc.font = Font(size=10)
            if ot == "Summa (alv 0)":
                cc.number_format = '#,##0.00'
            if ot == "Y-tunnus" and not rv.get(ot):
                cc.fill = PatternFill("solid", fgColor="FFCDD2")
    # Yhteensä
    yht_r = r0 + 1 + len(df)
    ws.cell(yht_r, 3, "YHTEENSÄ").font = Font(bold=True)
    tot = ws.cell(yht_r, 4, round(float(df["Summa (alv 0)"].sum()), 2) if not df.empty else 0)
    tot.font = Font(bold=True); tot.number_format = '#,##0.00'

    for col, w in zip("ABCDE", [30, 14, 22, 15, 9]):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def urakkailmoitus_csv(df: pd.DataFrame) -> bytes:
    return df.to_csv(index=False, sep=";").encode("utf-8-sig")
